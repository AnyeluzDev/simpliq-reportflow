import io
from pathlib import Path
import pandas as pd
import streamlit as st
import plotly.express as px
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

st.set_page_config(page_title="ReportFlow | SimpliQ", page_icon="logo_simpliq.png", layout="wide", initial_sidebar_state="expanded")
st.markdown("""<style>
@import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700;800&display=swap');
html,body,[class*="css"],.stApp{font-family:Inter,Arial,sans-serif}
.stApp,[data-testid="stAppViewContainer"]{background:#06172a;color:#e9f5ff}
[data-testid="stHeader"]{background:transparent}
[data-testid="stSidebar"]{background:#071f38;border-right:1px solid #173954}
.block-container{padding-top:1.4rem;max-width:1600px}
h1,h2,h3{color:#f1f8ff!important;letter-spacing:-.035em}
h1{font-weight:800!important}h3{font-size:1.12rem!important}
[data-testid="stMetric"]{background:linear-gradient(135deg,#102b49,#0b223c);border:1px solid #21415e;border-radius:16px;padding:19px 20px;min-height:124px}
[data-testid="stMetricLabel"]{color:#aac4dc}[data-testid="stMetricValue"]{color:white;font-weight:750;font-size:1.8rem}
[data-testid="stFileUploader"]{border:1px dashed #2086cc;border-radius:15px;background:#0b2744;padding:12px}
[data-testid="stFileUploaderDropzone"]{background:transparent;border:0}
div.stButton>button,div.stDownloadButton>button{background:linear-gradient(90deg,#0cb7ef,#1260fa);color:white;border:0;border-radius:10px;font-weight:700;min-height:42px}
div.stButton>button:hover,div.stDownloadButton>button:hover{color:white;border:1px solid #80d9ff}
[data-testid="stExpander"]{background:#0b223b;border:1px solid #24425d;border-radius:12px}
[data-testid="stDataFrame"]{border:1px solid #23435e;border-radius:10px}
[data-testid="stAlert"]{border-radius:12px}
.hero{background:linear-gradient(110deg,#0c2b4a,#09223c 70%,#0d315a);border:1px solid #1a4565;border-radius:18px;padding:25px 28px;margin-bottom:18px}
.eyebrow{font-size:12px;letter-spacing:.12em;color:#44ccff;font-weight:800;text-transform:uppercase}
.hero-title{font-size:clamp(26px,3vw,43px);font-weight:800;color:#fff;letter-spacing:-.04em;margin:7px 0}
.hero-copy{color:#aec9dd;font-size:14px;line-height:1.6}
.brand{text-align:center;padding:12px 0 20px}.brand img{width:88px;max-width:70%;display:block;margin:0 auto 5px}
.brand-title{font-weight:800;color:#f3f9ff;font-size:24px}.brand-sub{color:#9dbbd3;font-size:12px}
.section-caption{color:#9ab9cf;font-size:13px;margin-bottom:8px}
.footer-note{color:#9bb6ca;font-size:12px;margin-top:25px;text-align:center}
@media(max-width:750px){.block-container{padding:1rem}.hero{padding:18px}}
</style>""",unsafe_allow_html=True)
import base64
logo_b64=base64.b64encode((Path(__file__).parent/'logo_simpliq.png').read_bytes()).decode()
with st.sidebar:
    st.markdown(f'<div class="brand"><img src="data:image/png;base64,{logo_b64}"/><div class="brand-title">SimpliQ</div><div class="brand-sub">Convertimos horas de trabajo en minutos</div></div>',unsafe_allow_html=True)
    st.markdown('### 🧭 Navegación')
    st.markdown("🏠 **Resumen**")
    st.markdown('**📤 Importación de archivos**')
    st.markdown('**🛡️ Validaciones**')
    st.markdown('**📊 Análisis y gráficas**')
    st.markdown('**📋 Detalle de datos**')
    st.markdown('**📥 Exportación**')
    st.divider()
    st.info('💡 Puedes probar con nuestro CSV ficticio o cargar tu propio archivo de prueba.')
    st.caption('ReportFlow · Demo 01 · [simpliq.mx](https://simpliq.mx)')
st.markdown('<div class="hero"><div class="eyebrow">● Demo interactiva · Automatización de reportes</div><div class="hero-title">SimpliQ <span style="color:#34c5fa">ReportFlow</span></div><div class="hero-copy">De archivos desordenados a reportes profesionales.<br/>Importa ventas, detecta errores y descarga indicadores listos para tomar decisiones.</div></div>',unsafe_allow_html=True)
st.caption('🔒 Demo pública: los archivos se procesan durante la sesión; no se guardan en una base de datos. Evita subir información personal o confidencial.')
REQUIRED=['Fecha','Folio','Vendedor','Producto','Cantidad','Precio unitario','Sucursal','Estado']
SAMPLE=Path(__file__).parent/'data'/'ventas_demo.csv'
with open(SAMPLE,'rb') as f:
    sample_bytes=f.read()
st.subheader('📤 Cargar archivo de ventas')
left,right=st.columns([1.3,1],gap='large')
with left:
    file=st.file_uploader('Arrastra tu CSV o Excel (.xlsx)',type=['csv','xlsx'],help='Máximo 15 MB por archivo.')
with right:
    st.markdown('**¿No tienes un archivo?**')
    st.markdown('<div class="section-caption">Descarga nuestra muestra con 1,000 ventas ficticias y errores intencionales para comprobar las validaciones.</div>',unsafe_allow_html=True)
    st.download_button('⬇ Descargar CSV de demostración',sample_bytes,'ventas_demo.csv','text/csv')
use_demo=st.checkbox('Usar directamente los datos de demostración',value=file is None)
if file is None and not use_demo:
    st.stop()
try:
    if file is None:
        raw=pd.read_csv(io.BytesIO(sample_bytes),dtype={'Folio':str})
    elif file.name.lower().endswith('.xlsx'):
        raw=pd.read_excel(file,dtype={'Folio':str})
    else:
        blob=file.getvalue()
        try: raw=pd.read_csv(io.BytesIO(blob),encoding='utf-8-sig',sep=None,engine='python',dtype={'Folio':str})
        except UnicodeDecodeError: raw=pd.read_csv(io.BytesIO(blob),encoding='latin-1',sep=None,engine='python',dtype={'Folio':str})
except Exception as exc:
    st.error(f'No se pudo leer el archivo: {exc}')
    st.stop()
raw.columns=[str(c).strip() for c in raw.columns]
missing=[c for c in REQUIRED if c not in raw.columns]
if missing:
    st.error('Faltan columnas obligatorias: '+', '.join(missing))
    st.code(' | '.join(REQUIRED))
    st.stop()
if len(raw)>100_000:
    st.error('Esta demo admite hasta 100,000 filas por archivo.')
    st.stop()

df=raw[REQUIRED].copy()
for c in ['Folio','Vendedor','Producto','Sucursal','Estado']:
    df[c]=df[c].astype('string').str.strip()
df['Fecha']=pd.to_datetime(df['Fecha'],dayfirst=True,errors='coerce')
for c in ['Cantidad','Precio unitario']:
    df[c]=pd.to_numeric(df[c],errors='coerce')
status=df['Estado'].str.casefold().fillna('')
duplicates=df['Folio'].duplicated(keep='first') & df['Folio'].notna() & df['Folio'].ne('')
problems=pd.DataFrame(index=df.index)
problems['Motivo']=''
def add_problem(mask,label):
    problems.loc[mask,'Motivo']=problems.loc[mask,'Motivo'].apply(lambda s: (s+'; ' if s else '')+label)
add_problem(df['Fecha'].isna(),'Fecha inválida')
add_problem(df['Folio'].isna() | df['Folio'].eq(''),'Folio vacío')
add_problem(duplicates,'Folio duplicado')
add_problem(df[['Vendedor','Producto','Sucursal']].isna().any(axis=1) | df[['Vendedor','Producto','Sucursal']].eq('').any(axis=1),'Campo de texto vacío')
add_problem(df['Cantidad'].isna() | (df['Cantidad']<=0) | (df['Cantidad']%1!=0),'Cantidad inválida')
add_problem(df['Precio unitario'].isna() | (df['Precio unitario']<0),'Precio inválido')
add_problem(~status.isin(['completada','cancelada','pendiente']),'Estado desconocido')
invalid=problems['Motivo'].ne('')
review=df[invalid].copy()
review.insert(len(review.columns),'Motivo de revisión',problems.loc[invalid,'Motivo'])
valid=df[~invalid].copy()
valid['Importe']=valid['Cantidad']*valid['Precio unitario']
completed=valid[valid['Estado'].str.casefold()=='completada'].copy()

st.subheader('📊 Resumen de procesamiento')
a,b,c,d=st.columns(4)
a.metric('Registros recibidos',f'{len(df):,}')
b.metric('Registros válidos',f'{len(valid):,}')
c.metric('Para revisión',f'{len(review):,}')
d.metric('Ventas completadas',f'${completed["Importe"].sum():,.2f}')
st.caption('Las ventas incluyen únicamente operaciones válidas con estado Completada; se excluyen pendientes y canceladas.')
if len(completed):
    by_vendor=completed.groupby('Vendedor',as_index=False)['Importe'].sum().sort_values('Importe',ascending=False)
    by_branch=completed.groupby('Sucursal',as_index=False)['Importe'].sum().sort_values('Importe',ascending=False)
    by_product=completed.groupby('Producto',as_index=False)['Cantidad'].sum().sort_values('Cantidad',ascending=False)
    chart1,chart2,chart3=st.columns([1.2,1.2,1],gap='medium')
    def polish(fig):
        fig.update_layout(paper_bgcolor='rgba(0,0,0,0)',plot_bgcolor='rgba(0,0,0,0)',font_color='#c8deef',margin=dict(l=6,r=6,t=44,b=8),height=330,showlegend=False,title_font_size=15)
        fig.update_xaxes(showgrid=False,zeroline=False)
        fig.update_yaxes(gridcolor='#1c3a56',zeroline=False)
        return fig
    with chart1:
        fig=px.bar(by_vendor,x='Vendedor',y='Importe',title='Ventas por vendedor',color_discrete_sequence=['#1498fa'])
        st.plotly_chart(polish(fig),use_container_width=True)
    with chart2:
        fig=px.bar(by_branch,x='Sucursal',y='Importe',title='Ventas por sucursal',color_discrete_sequence=['#24c6df'])
        st.plotly_chart(polish(fig),use_container_width=True)
    with chart3:
        fig=px.pie(by_product,names='Producto',values='Cantidad',hole=.62,title='Productos vendidos',color_discrete_sequence=['#176bff','#0ac4cc','#fa4ba3','#7a5cff','#f7a44c','#5bd4fc'])
        fig=polish(fig);fig.update_layout(showlegend=True,legend=dict(orientation='h',y=-.18,x=0,font_size=10))
        st.plotly_chart(fig,use_container_width=True)
    x,y=st.columns(2)
    x.metric('Ticket promedio',f'${completed["Importe"].mean():,.2f}')
    y.metric('Unidades vendidas',f'{int(completed["Cantidad"].sum()):,}')
else:
    st.warning('No hay ventas completadas válidas para graficar.')
st.subheader('🛡️ Validaciones y detalle de datos')
with st.expander(f'🔎 Registros para revisión ({len(review)})',expanded=bool(len(review))):
    st.dataframe(review,use_container_width=True,hide_index=True)
with st.expander(f'✔️ Registros válidos ({len(valid)})'):
    st.dataframe(valid,use_container_width=True,hide_index=True)
def workbook_bytes():
    wb=Workbook()
    summary=wb.active;summary.title='Resumen'
    summary.append(['SIMPLIQ · REPORTFLOW',''])
    for k,v in [('Registros recibidos',len(df)),('Registros válidos',len(valid)),('Registros para revisión',len(review)),('Operaciones completadas',len(completed)),('Ventas completadas',float(completed['Importe'].sum())),('Ticket promedio',float(completed['Importe'].mean()) if len(completed) else 0)]:summary.append([k,v])
    for name,frame in [('Ventas válidas',valid),('Revisión',review),('Por vendedor',completed.groupby('Vendedor',as_index=False)['Importe'].sum() if len(completed) else pd.DataFrame(columns=['Vendedor','Importe']))]:
        ws=wb.create_sheet(name)
        ws.append(list(frame.columns))
        for rec in frame.itertuples(index=False,name=None):
            ws.append([v.to_pydatetime() if isinstance(v,pd.Timestamp) and pd.notna(v) else (None if pd.isna(v) else v) for v in rec])
    for ws in wb:
        ws.freeze_panes='A2'
        ws.auto_filter.ref=ws.dimensions if ws!=summary else 'A1:B7'
        for cell in ws[1]:
            cell.fill=PatternFill('solid',fgColor='103D58');cell.font=Font(color='FFFFFF',bold=True)
        for col in ws.columns:
            letter=get_column_letter(col[0].column)
            width=min(45,max(14,max(len(str(c.value or '')) for c in list(col)[:150])+2))
            ws.column_dimensions[letter].width=width
        if ws.title in ['Ventas válidas','Revisión']:
            for cell in ws['A'][1:]:cell.number_format='DD/MM/YYYY'
    summary['B6'].number_format='"$"#,##0.00';summary['B7'].number_format='"$"#,##0.00'
    out=io.BytesIO();wb.save(out);return out.getvalue()

st.subheader('📥 Exportar reporte')
st.markdown('Descarga un Excel con **resumen, ventas válidas, registros para revisión y resultados por vendedor**.')
st.download_button('📥 Descargar reporte Excel',workbook_bytes(),'SimpliQ_ReportFlow_Resultado.xlsx',mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',type='primary')
st.markdown('<div class="footer-note">SimpliQ ReportFlow · Demo 01 · <a href="https://simpliq.mx" style="color:#35c9ff">simpliq.mx</a> · Datos ficticios</div>',unsafe_allow_html=True)
